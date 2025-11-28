from pptx import Presentation
from pptx.util import Inches, Pt
import os

class PPTGenerator:
    def __init__(self):
        pass

    def generate_ppt(self, slides_data, images_folder, output_path):
        """
        Generates a PowerPoint presentation based on slides_data and images.
        """
        prs = Presentation()
        
        # Set Slide Size to 16:9 (13.333 in x 7.5 in)
        prs.slide_width = Inches(13.333)
        prs.slide_height = Inches(7.5)
        
        for slide_info in slides_data['slides']:
            layout_type = slide_info.get('layout', 'Title_Only')
            
            # Using Blank Layout (6) for maximum control, or Title Only (5)
            slide_layout = prs.slide_layouts[5] 
            slide = prs.slides.add_slide(slide_layout)
            
            # Title Styling
            title = slide.shapes.title
            title.text = slide_info.get('title', 'No Title')
            title.left = Inches(0.5)
            title.top = Inches(0.5)
            title.width = Inches(12.333)
            title.height = Inches(1.0)
            
            if layout_type == 'Title_Only':
                # Full width text
                left = Inches(1)
                top = Inches(2)
                width = Inches(11.333)
                height = Inches(5)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.word_wrap = True
                tf.text = slide_info.get('content', '')

            elif layout_type == 'Title_Image_Right':
                # Content on Left
                left = Inches(0.5)
                top = Inches(2)
                width = Inches(6)
                height = Inches(5)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.word_wrap = True
                tf.text = slide_info.get('content', '')
                
                # Image on Right
                img_path = os.path.join(images_folder, f"slide_{slide_info['id']}.png")
                if os.path.exists(img_path):
                    left = Inches(7)
                    top = Inches(2)
                    width = Inches(5.8)
                    # height will be auto-scaled to maintain aspect ratio if width is given, 
                    # but we can constrain it.
                    slide.shapes.add_picture(img_path, left, top, width=width)

            elif layout_type == 'Title_Image_Left':
                # Image on Left
                img_path = os.path.join(images_folder, f"slide_{slide_info['id']}.png")
                if os.path.exists(img_path):
                    left = Inches(0.5)
                    top = Inches(2)
                    width = Inches(5.8)
                    slide.shapes.add_picture(img_path, left, top, width=width)

                # Content on Right
                left = Inches(6.8)
                top = Inches(2)
                width = Inches(6)
                height = Inches(5)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.word_wrap = True
                tf.text = slide_info.get('content', '')

            elif layout_type == 'Two_Column':
                # Column 1
                left = Inches(0.5)
                top = Inches(2)
                width = Inches(6)
                height = Inches(5)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.word_wrap = True
                tf.text = slide_info.get('content', '')
                
                # Column 2
                left = Inches(6.8)
                txBox2 = slide.shapes.add_textbox(left, top, width, height)
                tf2 = txBox2.text_frame
                tf2.word_wrap = True
                tf2.text = " " # Placeholder

            else:
                # Default fallback
                left = Inches(1)
                top = Inches(2)
                width = Inches(11.333)
                height = Inches(5)
                txBox = slide.shapes.add_textbox(left, top, width, height)
                tf = txBox.text_frame
                tf.text = slide_info.get('content', '')

        prs.save(output_path)
        return output_path
